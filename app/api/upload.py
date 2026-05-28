import io

import docx
import openpyxl
import pdfplumber
from fastapi import APIRouter, File, HTTPException, UploadFile
from openai import OpenAI
from pydantic import BaseModel

from app.config import settings
from app.db import insert_chunks, similarity_search

router = APIRouter(prefix="/documents", tags=["documents"])

CHUNK_SIZE = 500  # characters per chunk
CHUNK_OVERLAP = 50


# ── helpers ──────────────────────────────────────────────────────────────────

def _extract_text(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n".join(p.extract_text() or "" for p in pdf.pages)
    if ext == "docx":
        doc = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs)
    if ext in ("xlsx", "xls"):
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                line = "\t".join("" if v is None else str(v) for v in row)
                if line.strip():
                    lines.append(line)
        return "\n".join(lines)
    raise HTTPException(status_code=400, detail=f"Unsupported file type: .{ext}")


def _chunk_text(text: str) -> list[str]:
    words = text.split()
    full = " ".join(words)
    chunks, start = [], 0
    while start < len(full):
        end = start + CHUNK_SIZE
        chunks.append(full[start:end])
        start += CHUNK_SIZE - CHUNK_OVERLAP
    return [c for c in chunks if c.strip()]


def _embed(texts: list[str]) -> list[list[float]]:
    if not settings.openai_api_key:
        raise HTTPException(status_code=500, detail="OpenAI API key not configured")
    client = OpenAI(api_key=settings.openai_api_key)
    resp = client.embeddings.create(model=settings.openai_embedding_model, input=texts)
    return [item.embedding for item in resp.data]


# ── endpoints ─────────────────────────────────────────────────────────────────

@router.post("/upload", summary="Upload a file and store embeddings")
async def upload_file(file: UploadFile = File(...)):
    data = await file.read()
    text = _extract_text(file.filename, data)
    if not text.strip():
        raise HTTPException(status_code=422, detail="No text could be extracted from the file")

    chunks = _chunk_text(text)
    embeddings = _embed(chunks)

    records = [
        {"chunk_index": i, "content": c, "embedding": e}
        for i, (c, e) in enumerate(zip(chunks, embeddings))
    ]
    insert_chunks(file.filename, records)

    return {"filename": file.filename, "chunks_stored": len(records)}


class QueryRequest(BaseModel):
    question: str
    top_k: int = 5


@router.post("/query", summary="Query documents and get an LLM answer")
async def query_documents(body: QueryRequest):
    if not settings.openai_api_key:
        raise HTTPException(status_code=500, detail="OpenAI API key not configured")

    # Embed the question
    q_embedding = _embed([body.question])[0]

    # Retrieve relevant chunks
    hits = similarity_search(q_embedding, top_k=body.top_k)
    if not hits:
        raise HTTPException(status_code=404, detail="No relevant chunks found")

    context = "\n\n".join(
        f"[{h['filename']} chunk {h['chunk_index']}]\n{h['content']}" for h in hits
    )

    # Ask OpenAI
    client = OpenAI(api_key=settings.openai_api_key)
    completion = client.chat.completions.create(
        model=settings.openai_model,
        messages=[
            {
                "role": "system",
                "content": "You are a helpful assistant. Answer using only the provided context.",
            },
            {
                "role": "user",
                "content": f"Context:\n{context}\n\nQuestion: {body.question}",
            },
        ],
    )
    answer = completion.choices[0].message.content

    return {
        "answer": answer,
        "sources": [{"filename": h["filename"], "chunk_index": h["chunk_index"], "score": h["score"]} for h in hits],
    }

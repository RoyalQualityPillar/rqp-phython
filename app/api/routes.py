import io
import zipfile
from datetime import datetime, timezone

import docx
import openpyxl
import pdfplumber
from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from app.config import settings
from app.db import insert_chunks, similarity_search
from app.llm import chat, embed
from app.logger import ERROR_LOG, INFO_LOG, WARNING_LOG, get_logger

log = get_logger("routes")
router = APIRouter()

CHUNK_SIZE = 500
CHUNK_OVERLAP = 50


# ── helpers ───────────────────────────────────────────────────────────────────

def _extract_text(filename: str, data: bytes) -> str:
    ext = filename.rsplit(".", 1)[-1].lower()
    log.info("Extracting text from '%s' (type=%s, size=%d bytes)", filename, ext, len(data))
    if ext == "pdf":
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            text = "\n".join(p.extract_text() or "" for p in pdf.pages)
    elif ext == "docx":
        doc = docx.Document(io.BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
    elif ext in ("xlsx", "xls"):
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                line = "\t".join("" if v is None else str(v) for v in row)
                if line.strip():
                    lines.append(line)
        text = "\n".join(lines)
    else:
        log.warning("Unsupported file type requested: .%s", ext)
        raise HTTPException(status_code=400, detail=f"Unsupported file type: .{ext}")
    log.info("Extracted %d characters from '%s'", len(text), filename)
    return text


def _chunk_text(text: str) -> list[str]:
    full = " ".join(text.split())
    chunks, start = [], 0
    while start < len(full):
        chunks.append(full[start: start + CHUNK_SIZE])
        start += CHUNK_SIZE - CHUNK_OVERLAP
    log.info("Text split into %d chunks (size=%d, overlap=%d)", len(chunks), CHUNK_SIZE, CHUNK_OVERLAP)
    return [c for c in chunks if c.strip()]


# ── health ────────────────────────────────────────────────────────────────────

@router.get("/health", tags=["health"], summary="Health check")
async def health_check() -> dict:
    log.info("Health check called")
    return {
        "status": "ok",
        "service": "rqp-chatbot-backend",
        "time": datetime.now(timezone.utc).isoformat(),
    }


@router.get("/live", tags=["health"], summary="Liveness probe")
async def live() -> dict:
    return {"status": "live"}


@router.get("/ready", tags=["health"], summary="Readiness probe")
async def ready() -> dict:
    return {"status": "ready"}


# ── chatbot ───────────────────────────────────────────────────────────────────

class ChatRequest(BaseModel):
    message: str


@router.post("/chatbot/dummy", tags=["chatbot"], summary="Dummy chatbot endpoint")
async def chatbot_dummy(payload: dict) -> dict:
    log.info("Dummy chatbot called — payload_keys=%s", list(payload.keys()))
    reply_text = "This is a dummy response from the chatbot backend."
    llm_used = False
    llm_error: str | None = None

    if "message" in payload:
        try:
            reply_text = chat(
                system="You are a helpful assistant for the RQP chatbot UI.",
                user=str(payload["message"]),
            )
            llm_used = True
            log.info("Dummy chatbot LLM reply generated successfully")
        except Exception as exc:
            llm_error = str(exc)
            log.warning("Dummy chatbot LLM call failed, using fallback reply: %s", exc)

    return {
        "reply": reply_text,
        "received": payload,
        "provider": settings.llm_provider,
        "llm_used": llm_used,
        "llm_error": llm_error,
        "timestamp": datetime.now(timezone.utc).isoformat(),
    }


@router.post("/chatbot/openai", tags=["chatbot"], summary="Chat via LLM (Gemini or OpenAI)")
async def chatbot_openai(body: ChatRequest) -> dict:
    log.info("Chat endpoint called — provider=%s msg_len=%d", settings.llm_provider, len(body.message))
    try:
        reply = chat(
            system="You are a helpful assistant for the RQP chatbot UI.",
            user=body.message,
        )
        log.info("Chat reply generated successfully")
    except Exception as exc:
        log.error("Chat endpoint error: %s", exc, exc_info=True)
        raise HTTPException(status_code=500, detail=str(exc)) from exc
    return {
        "reply": reply,
        "provider": settings.llm_provider,
        "timestamp": datetime.now(timezone.utc).isoformat(),
    }


# ── documents ─────────────────────────────────────────────────────────────────

@router.post("/documents/upload", tags=["documents"], summary="Upload a file and store embeddings")
async def upload_file(file: UploadFile = File(...)):
    log.info("Upload request — filename='%s' content_type=%s", file.filename, file.content_type)
    data = await file.read()
    log.info("File read — size=%d bytes", len(data))

    text = _extract_text(file.filename, data)
    if not text.strip():
        log.warning("No text extracted from '%s'", file.filename)
        raise HTTPException(status_code=422, detail="No text could be extracted from the file")

    chunks = _chunk_text(text)
    log.info("Generating embeddings for %d chunks", len(chunks))
    embeddings = embed(chunks)

    records = [
        {"chunk_index": i, "content": c, "embedding": e}
        for i, (c, e) in enumerate(zip(chunks, embeddings))
    ]
    insert_chunks(file.filename, records)
    log.info("Upload complete — '%s' stored %d chunks", file.filename, len(records))
    return {"filename": file.filename, "chunks_stored": len(records)}


class QueryRequest(BaseModel):
    question: str
    top_k: int = 5


@router.post("/documents/query", tags=["documents"], summary="Query documents and get an LLM answer")
async def query_documents(body: QueryRequest):
    log.info("Query request — question_len=%d top_k=%d", len(body.question), body.top_k)

    q_embedding = embed([body.question])[0]
    hits = similarity_search(q_embedding, top_k=body.top_k)

    if not hits:
        log.warning("No relevant chunks found for query: '%s'", body.question[:80])
        raise HTTPException(status_code=404, detail="No relevant chunks found")

    log.info("Found %d relevant chunks, sending to LLM", len(hits))
    context = "\n\n".join(
        f"[{h['filename']} chunk {h['chunk_index']}]\n{h['content']}" for h in hits
    )
    answer = chat(
        system="You are a helpful assistant. Answer using only the provided context.",
        user=f"Context:\n{context}\n\nQuestion: {body.question}",
    )
    log.info("Query answered successfully — answer_len=%d", len(answer))
    return {
        "answer": answer,
        "provider": settings.llm_provider,
        "sources": [
            {"filename": h["filename"], "chunk_index": h["chunk_index"], "score": h["score"]}
            for h in hits
        ],
    }


# ── logs ──────────────────────────────────────────────────────────────────────

@router.get("/logs/download", tags=["logs"], summary="Download all log files as a ZIP")
async def download_logs():
    log.info("Log download requested")
    log_files = {
        "info.log": INFO_LOG,
        "warning.log": WARNING_LOG,
        "error.log": ERROR_LOG,
    }

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, mode="w", compression=zipfile.ZIP_DEFLATED) as zf:
        for name, path in log_files.items():
            if path.exists():
                zf.write(path, arcname=name)
                log.info("Added '%s' to zip (%d bytes)", name, path.stat().st_size)
            else:
                zf.writestr(name, "")
                log.warning("Log file '%s' not found, adding empty file", name)

    buf.seek(0)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"rqp_logs_{timestamp}.zip"
    log.info("Serving log zip: %s", filename)
    return StreamingResponse(
        buf,
        media_type="application/zip",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
